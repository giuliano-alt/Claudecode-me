"""Genera utenza.js dal resoconto orientamento."""
import openpyxl, re, datetime, json

P = '/root/.claude/uploads/a01855de-95c7-57ca-b448-f7b3255ad6fc/8b45284d-Resoconto_orientamento_aggiornato_al_27.07.xlsx'
ws = openpyxl.load_workbook(P, data_only=True)['Foglio1']
DATE_RE = re.compile(r'(\d{2})[/.](\d{2})[/.](\d{4})')

def parse(v):
    if v is None: return []
    if isinstance(v, datetime.datetime): return [v.date()]
    s = str(v); out = []
    for d, m, y in DATE_RE.findall(s):
        try: out.append(datetime.date(int(y), int(m), int(d)))
        except ValueError: pass
    if not out:
        m = re.search(r'(\d{2})/(\d{2})(\d{4})', s)
        if m:
            try: out.append(datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
            except ValueError: pass
    return out

BORSISTI = {'MARONG Ablie','CAMARA Ousmane','KRUBALLY Sainey','TRAORE Baba','MANNEH Abubacarr',
            'TALEEM Usama','ISLAM Md Jahidul','HAWLADER Omar','VINCENT Gift'}
BORSA_NOTE = {
    'VINCENT Gift': 'Borsa + tirocinio',
    'TRAORE Baba': 'Borsa conclusa',
    'CAMARA Ousmane': 'Borsa conclusa',
    'MARONG Ablie': 'Borsa in corso',
    'MANNEH Abubacarr': 'Borsa in corso',
    'KRUBALLY Sainey': 'Borsa in corso',
    'HAWLADER Omar': 'Borsa in corso',
    'TALEEM Usama': 'Borsa interrotta',
    'ISLAM Md Jahidul': 'Borsa non avviata',
}

fmt = lambda d: d.strftime('%d/%m/%Y')

def cell(ds):
    if not ds: return '—'
    s = fmt(min(ds))
    if len(ds) > 1: s += f' (+{len(ds)-1})'
    return s

rows, n = [], 0
for r in list(range(3, 51)) + list(range(56, 60)):
    nome = ws.cell(r, 3).value
    if not nome or str(nome).strip().upper().startswith('COGNOME'): continue
    nome = str(nome).strip()
    n += 1
    ing, fine = parse(ws.cell(r, 4).value), parse(ws.cell(r, 5).value)
    col, cv, cmp_ = parse(ws.cell(r, 6).value), parse(ws.cell(r, 7).value), parse(ws.cell(r, 8).value)
    rows.append({
        'n': n,
        'prot': str(ws.cell(r, 2).value),
        'nome': nome,
        'tipo': 'N' if r >= 56 else 'S',
        'ing': fmt(ing[0]) if ing else '—',
        'fine': fmt(fine[0]) if fine else 'n.d.',
        'col': cell(col), 'cv': cell(cv), 'comp': cell(cmp_),
        'borsa': BORSA_NOTE.get(nome, '—'),
    })

with open('utenza.js', 'w') as f:
    f.write('/* Generato da gen_utenza.py — fonte: Resoconto orientamento al 27/07/2026 */\n')
    f.write('module.exports = ' + json.dumps(rows, ensure_ascii=False, indent=0).replace('\n', '\n') + ';\n')

print('record scritti:', len(rows))
print('borsisti riconosciuti:', sum(1 for x in rows if x['borsa'] != '—'))
mx = max(len(x['nome']) for x in rows)
print('nome piu lungo:', mx, [x['nome'] for x in rows if len(x['nome']) == mx])
