import openpyxl, re, datetime, json
from collections import Counter

P = '/root/.claude/uploads/a01855de-95c7-57ca-b448-f7b3255ad6fc/8b45284d-Resoconto_orientamento_aggiornato_al_27.07.xlsx'
ws = openpyxl.load_workbook(P, data_only=True)['Foglio1']

HEADER_ROWS = {2, 16, 41, 55}
SINGLE_ROWS = list(range(3, 51))
NUCLEI_ROWS = list(range(56, 60))

DATE_RE = re.compile(r'(\d{2})[/.](\d{2})[/.](\d{4})')

def parse_cell(v):
    """Restituisce la lista di date contenute nella cella (iniziale + aggiornamenti)."""
    if v is None:
        return []
    if isinstance(v, datetime.datetime):
        return [v.date()]
    s = str(v)
    out = []
    for d, m, y in DATE_RE.findall(s):
        try:
            out.append(datetime.date(int(y), int(m), int(d)))
        except ValueError:
            pass
    # date malformate tipo '22/072025'
    if not out:
        m = re.search(r'(\d{2})/(\d{2})(\d{4})', s)
        if m:
            try:
                out.append(datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
            except ValueError:
                pass
    return out

recs = []
for r in SINGLE_ROWS + NUCLEI_ROWS:
    if r in HEADER_ROWS:
        continue
    nome = ws.cell(r, 3).value
    if not nome or str(nome).strip().upper().startswith('COGNOME'):
        continue
    recs.append({
        'row': r,
        'tipo': 'nucleo' if r in NUCLEI_ROWS else 'singolo',
        'prot': ws.cell(r, 2).value,
        'nome': str(nome).strip(),
        'ingresso': parse_cell(ws.cell(r, 4).value),
        'fine': parse_cell(ws.cell(r, 5).value),
        'colloquio': parse_cell(ws.cell(r, 6).value),
        'cv': parse_cell(ws.cell(r, 7).value),
        'comp': parse_cell(ws.cell(r, 8).value),
        'raw_ing': ws.cell(r, 4).value,
    })

n = len(recs)
sing = [x for x in recs if x['tipo'] == 'singolo']
nucl = [x for x in recs if x['tipo'] == 'nucleo']
print(f'BENEFICIARI IN CARICO: {n}  (singoli {len(sing)}, nuclei {len(nucl)})')

con_col = [x for x in recs if x['colloquio']]
con_cv  = [x for x in recs if x['cv']]
con_cmp = [x for x in recs if x['comp']]
print(f'\nCOPERTURA')
print(f'  colloquio orientamento : {len(con_col):2d}/{n}  {len(con_col)/n*100:.1f}%')
print(f'  curriculum vitae       : {len(con_cv):2d}/{n}  {len(con_cv)/n*100:.1f}%')
print(f'  determinazione compet. : {len(con_cmp):2d}/{n}  {len(con_cmp)/n*100:.1f}%')

tot_col = sum(len(x['colloquio']) for x in recs)
tot_cv  = sum(len(x['cv']) for x in recs)
tot_cmp = sum(len(x['comp']) for x in recs)
print(f'\nAZIONI REGISTRATE (iniziali + aggiornamenti)')
print(f'  colloqui        : {tot_col}   (di cui {tot_col-len(con_col)} aggiornamenti)')
print(f'  CV redatti/agg. : {tot_cv}   (di cui {tot_cv-len(con_cv)} aggiornamenti)')
print(f'  determinazioni  : {tot_cmp}')
print(f'  TOTALE AZIONI   : {tot_col+tot_cv+tot_cmp}')

# azioni per anno
per_anno = Counter()
for x in recs:
    for k in ('colloquio', 'cv', 'comp'):
        for d in x[k]:
            per_anno[d.year] += 1
print(f'\nAZIONI PER ANNO: {dict(sorted(per_anno.items()))}')

per_anno_k = {}
for k in ('colloquio', 'cv', 'comp'):
    c = Counter()
    for x in recs:
        for d in x[k]:
            c[d.year] += 1
    per_anno_k[k] = dict(sorted(c.items()))
print('  dettaglio:', json.dumps(per_anno_k))

# 2026 fino al 27/07
a2026 = [d for x in recs for k in ('colloquio','cv','comp') for d in x[k] if d.year == 2026]
print(f'\n2026: {len(a2026)} azioni; ultima {max(a2026)}; prima {min(a2026)}')
mesi26 = Counter(d.month for d in a2026)
print('  per mese 2026:', dict(sorted(mesi26.items())))

# nessun contatto
zero = [x for x in recs if not x['colloquio'] and not x['cv'] and not x['comp']]
print(f'\nSENZA ALCUNA AZIONE REGISTRATA: {len(zero)}')
for x in zero:
    ing = x['ingresso'][0] if x['ingresso'] else None
    print(f"  r{x['row']:2d} prot {x['prot']:>4} {x['nome'][:42]:42} ingresso {ing}")

# tempi di presa in carico
lag = []
for x in recs:
    if x['ingresso'] and x['colloquio']:
        d = (min(x['colloquio']) - x['ingresso'][0]).days
        lag.append((d, x['nome'], x['ingresso'][0], min(x['colloquio'])))
lag.sort()
gg = [d for d, *_ in lag]
print(f'\nTEMPO INGRESSO -> PRIMO COLLOQUIO  (n={len(gg)})')
print(f'  min {min(gg)} gg | mediana {sorted(gg)[len(gg)//2]} gg | media {sum(gg)/len(gg):.0f} gg | max {max(gg)} gg')
entro30 = sum(1 for d in gg if d <= 30); entro60 = sum(1 for d in gg if d <= 60); entro90 = sum(1 for d in gg if d <= 90)
print(f'  entro 30 gg: {entro30} ({entro30/len(gg)*100:.0f}%) | entro 60: {entro60} ({entro60/len(gg)*100:.0f}%) | entro 90: {entro90} ({entro90/len(gg)*100:.0f}%)')
print('  i 5 piu lunghi:')
for d, nome, i, c in lag[-5:]:
    print(f'    {d:4d} gg  {nome[:40]:40} {i} -> {c}')

# ingressi per anno
ing_anno = Counter(x['ingresso'][0].year for x in recs if x['ingresso'])
print(f'\nINGRESSI IN PROGETTO PER ANNO: {dict(sorted(ing_anno.items()))}')

# uscite previste
fine_anno = Counter(x['fine'][0].year for x in recs if x['fine'])
print(f'USCITE PREVISTE PER ANNO: {dict(sorted(fine_anno.items()))}')
scad26 = [(x['fine'][0], x['nome']) for x in recs if x['fine'] and x['fine'][0].year == 2026]
scad26.sort()
print(f'  in scadenza nel 2026 ({len(scad26)}):')
for d, nm in scad26:
    print(f'    {d}  {nm[:40]}')
senza_fine = [x['nome'] for x in recs if not x['fine']]
print(f'  senza data di fine progetto ({len(senza_fine)}): {senza_fine}')

# incrocio con le borse lavoro
BORSE = ['VINCENT Gift','TRAORE Baba','CAMARA Ousmane','MARONG','TALEEM Usama',
         'HAWLADER Omar','ISLAM Md','MANNEH Abubacarr','KRUBALLY Sainey']
print('\nINCROCIO CON I PERCORSI DI BORSA LAVORO')
found = []
for b in BORSE:
    hit = [x for x in recs if b.lower().split()[0] in x['nome'].lower() and
           (len(b.split()) < 2 or b.lower().split()[1][:4] in x['nome'].lower())]
    for h in hit:
        found.append(h['nome'])
        print(f"  {b:20} -> r{h['row']:2d} prot {h['prot']:>4} {h['nome'][:38]:38} colloquio {min(h['colloquio']) if h['colloquio'] else '-'}")
print(f'  beneficiari con borsa: {len(set(found))} su {n} in carico = {len(set(found))/n*100:.1f}%')

# anomalie
print('\nANOMALIE DI COMPILAZIONE')
for x in recs:
    if isinstance(x['raw_ing'], str) and not DATE_RE.search(x['raw_ing']):
        print(f"  r{x['row']}: data ingresso malformata -> {x['raw_ing']!r} ({x['nome']})")
for x in recs:
    for k in ('colloquio', 'cv'):
        ds = x[k]
        if len(ds) > 1 and ds != sorted(ds):
            print(f"  r{x['row']}: date {k} non in ordine -> {ds} ({x['nome']})")
for x in recs:
    if x['comp'] and not x['colloquio']:
        print(f"  r{x['row']}: competenze senza colloquio ({x['nome']})")
    if x['cv'] and not x['colloquio']:
        print(f"  r{x['row']}: CV senza colloquio ({x['nome']})")
